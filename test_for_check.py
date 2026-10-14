#!/usr/bin/env python3
"""for-check.csv: snapshots never raise, rows are long-format, summary.xlsx gets a second sheet."""

from __future__ import annotations

import csv
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import for_check as fc


class _Proc:
    def __init__(self, code=0, out="", err=""):
        self.returncode, self.stdout, self.stderr = code, out, err


class LsRemoteTests(unittest.TestCase):
    def test_keeps_heads_tags_and_pr_refs_only_and_token_not_in_argv(self):
        seen = {}

        def fake_run(cmd, **kw):
            seen["cmd"], seen["env"] = cmd, kw["env"]
            return _Proc(out="a1\trefs/heads/main\nb2\trefs/tags/v1\nc3\trefs/pull/7/head\nd4\trefs/notes/x\n")

        refs = fc.ls_remote_refs("https://h/o/r.git", "x-access-token", "SECRET", run=fake_run)
        self.assertEqual(refs, {"refs/heads/main": "a1", "refs/tags/v1": "b2", "refs/pull/7/head": "c3"})
        self.assertNotIn("SECRET", " ".join(seen["cmd"]))
        self.assertIn("http.extraHeader", seen["env"]["GIT_CONFIG_KEY_0"])

    def test_failure_redacts_token(self):
        with self.assertRaises(RuntimeError) as ctx:
            fc.ls_remote_refs("u", "x", "SECRET", run=lambda *a, **k: _Proc(2, err="bad SECRET here"))
        self.assertNotIn("SECRET", str(ctx.exception))


class SnapshotTests(unittest.TestCase):
    def test_gitlab_counts_use_x_total_and_blank_on_failure(self):
        def fake_get(url, headers, **kw):
            if "statistics=true" in url:
                return {"statistics": {"commit_count": 9}, "default_branch": "main", "wiki_enabled": True}, {}
            if "/labels" in url:
                return [], {"X-Total": "5"}
            raise RuntimeError("HTTP 403")

        with mock.patch.object(fc, "http_get_json", fake_get):
            c = fc.gitlab_counts("g/p", token="t", host="gitlab.com")
        self.assertEqual((c["commits"], c["labels"], c["has_wiki"]), (9, 5, "true"))
        self.assertEqual(c["collaborators"], "")  # unknown stays blank, never a guess
        self.assertTrue(any(e.startswith("collaborators:") for e in c["_errors"]))

    def test_github_count_from_link_header(self):
        link = '<https://api.github.com/x?per_page=1&page=294>; rel="last"'
        with mock.patch.object(fc, "http_get_json", lambda *a, **k: ([{}], {"Link": link})):
            self.assertEqual(fc._count("https://api.github.com/x", {}), 294)

    def test_snapshot_never_raises(self):
        with mock.patch.object(fc, "github_counts", side_effect=RuntimeError("boom")), \
                mock.patch.object(fc, "ls_remote_refs", side_effect=RuntimeError("nope")):
            snap = fc.snapshot("github", "o/r", token="t", host="github.com")
        self.assertEqual(snap["counts"], {})
        self.assertEqual(len(snap["errors"]), 2)


class CountTests(unittest.TestCase):
    def test_unknown_when_gitlab_omits_total_and_more_pages_exist(self):
        hdr = {"Link": '<https://g/x?page=2&per_page=1>; rel="next"'}
        with mock.patch.object(fc, "http_get_json", lambda *a, **k: ([{}], hdr)):
            self.assertEqual(fc._count("https://g/x", {}), "")

    def test_single_item_and_empty_lists_are_real_counts(self):
        with mock.patch.object(fc, "http_get_json", lambda *a, **k: ([{}], {})):
            self.assertEqual(fc._count("https://g/x", {}), 1)
        with mock.patch.object(fc, "http_get_json", lambda *a, **k: ([], {})):
            self.assertEqual(fc._count("https://g/x", {}), 0)


class MergedCountTests(unittest.TestCase):
    def test_gitlab_has_merged_only_count_beside_all(self):
        def fake_get(url, headers, **kw):
            if "state=merged" in url:
                return [{}], {"X-Total": "3"}
            if "merge_requests" in url:
                return [{}], {"X-Total": "10"}
            return [], {"X-Total": "0"}

        with mock.patch.object(fc, "http_get_json", fake_get):
            c = fc.gitlab_counts("g/p", token="t", host="gitlab.com", project_id=1)
        self.assertEqual((c["prs"], c["merged_prs"]), (10, 3))

    def test_light_snapshot_skips_the_extra_calls(self):
        calls = []
        with mock.patch.object(fc, "http_get_json", lambda url, h, **k: (calls.append(url), ([], {"X-Total": "0"}))[1]):
            c = fc.gitlab_counts("g/p", token="t", host="gitlab.com", project_id=1, light=True)
        self.assertEqual(c["collaborators"], "")
        self.assertFalse(any("members" in u or "labels" in u for u in calls))


class RedactTests(unittest.TestCase):
    def test_empty_token_does_not_corrupt_message(self):
        self.assertEqual(fc.redact("fatal: repo not found", ""), "fatal: repo not found")

    def test_header_forms_and_userinfo(self):
        import base64

        b64 = base64.b64encode(b"x:SECRET").decode()
        out = fc.redact(f"Authorization: Basic {b64} and https://u:SECRET@h/x PRIVATE-TOKEN: SECRET", "SECRET")
        self.assertNotIn("SECRET", out)
        self.assertNotIn(b64, out)


class RefTests(unittest.TestCase):
    REFS = {"refs/heads/acme-ticket-42": "a1", "refs/tags/v1": "b2", "refs/pull/7/head": "c3"}

    def test_default_hashes_names_but_keeps_sha_and_type(self):
        out = fc.mask_refs(self.REFS, "hash")
        self.assertEqual(sorted(out.values()), ["a1", "b2", "c3"])
        self.assertIn("refs/pull/7/head", out)
        self.assertFalse(any("acme" in k for k in out))
        self.assertTrue(all(k.startswith(("refs/heads/h-", "refs/tags/h-", "refs/pull/")) for k in out))
        self.assertEqual(fc.mask_refs(self.REFS, "none"), {})
        self.assertEqual(fc.mask_refs(self.REFS, "names"), self.REFS)

    def test_after_is_a_diff_of_before(self):
        before = {"refs": {"r/a": "1", "r/b": "2"}, "refs_ok": True, "errors": []}
        after = {"refs": {"r/a": "1", "r/b": "9", "r/c": "3"}, "refs_ok": True, "errors": []}
        self.assertEqual(fc._ref_diff(before, after)["refs"], {"r/b": "9", "r/c": "3"})
        gone = fc._ref_diff(before, {"refs": {"r/a": "1"}, "refs_ok": True, "errors": []})
        self.assertEqual(gone["refs"], {"r/b": "deleted"})

    def test_failed_after_listing_is_not_read_as_everything_deleted(self):
        before = {"refs": {"r/a": "1"}, "refs_ok": True, "errors": []}
        after = {"refs": {}, "refs_ok": False, "errors": ["refs: boom"]}
        self.assertEqual(fc._ref_diff(before, after)["refs"], {})
        self.assertNotIn("refs_mode_after", fc._ref_diff(before, after))

    def test_bitbucket_user_is_email_when_known(self):
        self.assertEqual(fc._remote_url("bitbucket", "w/r", "", "me@x.io")[1], "me@x.io")
        self.assertEqual(fc._remote_url("bitbucket", "w/r", "")[1], "x-token-auth")


class SaveAndStaleTests(unittest.TestCase):
    def test_save_failure_is_logged_and_leaves_an_error_row(self):
        with tempfile.TemporaryDirectory() as d:
            run = Path(d)
            api = run / "api" / "o__r"
            api.mkdir(parents=True)
            (run / "repos.json").write_text(json.dumps([{"full_name": "o/r"}]))
            log = mock.Mock()
            real = Path.write_text
            calls = []

            def flaky(self, *a, **k):
                calls.append(1)
                if len(calls) == 1:
                    raise OSError("disk hiccup")
                return real(self, *a, **k)

            with mock.patch.object(Path, "write_text", flaky):
                fc.save_repo(api, platform="github", host="h", started_at="t", before={}, after={},
                             captured={"repo": "o/r"}, log=log)
            log.warning.assert_called()
            rows = fc.build_rows(run, {"o/r"})
        self.assertTrue(any(r["section"] == "error" and r["key"] == "save" for r in rows))

    def test_repos_not_in_this_run_are_filtered(self):
        with tempfile.TemporaryDirectory() as d:
            run = Path(d)
            for slug, name in (("a__x", "a/x"), ("b__y", "b/y")):
                api = run / "api" / slug
                api.mkdir(parents=True)
                fc.save_repo(api, platform="github", host="h", started_at="t", before={}, after={},
                             captured={"repo": name})
            (run / "repos.json").write_text(json.dumps([{"full_name": "a/x"}]))
            path = fc.write_run_csv(run)
            repos = {r["repo"] for r in csv.DictReader(path.open())}
        self.assertEqual(repos, {"a/x"})


class WorkbookCellTests(unittest.TestCase):
    def _sheet_values(self, summary_rows, fc_rows):
        import openpyxl

        with tempfile.TemporaryDirectory() as d:
            run = Path(d)
            with (run / "summary.csv").open("w", newline="") as fh:
                csv.writer(fh).writerows(summary_rows)
            with (run / fc.FILE_NAME).open("w", newline="") as fh:
                csv.writer(fh).writerows(fc_rows)
            wb = openpyxl.load_workbook(fc.write_workbook(run / "summary.csv"))
            return wb, [[c.value for c in r] for r in wb["for-check"].iter_rows()]

    def test_control_chars_numbers_and_formulas(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")
        wb, vals = self._sheet_values(
            [["org", "repo", "loc"], ["o", "r", "120"]],
            [fc.COLUMNS, ["o/r", "error", "before", "bad\x00\x07 text"],
             ["o/r", "counts_before", "prs", "294"], ["o/r", "ref_before", "refs/heads/x", "1234567"],
             ["o/r", "error", "x", "=HYPERLINK(\"http://evil\")"]],
        )
        self.assertEqual(vals[1][3], "bad text")
        self.assertEqual(vals[2][3], 294)
        self.assertEqual(vals[3][3], "1234567")  # a sha stays text
        self.assertEqual(vals[4][3], '=HYPERLINK("http://evil")')
        self.assertEqual(wb["for-check"]["D5"].data_type, "s")
        self.assertEqual(wb["summary"]["C2"].value, 120)

    def test_row_cap(self):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            self.skipTest("openpyxl not installed")
        rows = [fc.COLUMNS] + [["o/r", "counts_before", f"k{i}", str(i)] for i in range(10)]
        with mock.patch.object(fc, "XLSX_MAX_ROWS", 5):
            _, vals = self._sheet_values([["org"], ["o"]], rows)
        self.assertEqual(len(vals), 5 + 1)
        self.assertIn("truncated", vals[-1][0])

    def test_skip_for_check_sheet(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        with tempfile.TemporaryDirectory() as d:
            run = Path(d)
            (run / "summary.csv").write_text("org\no\n")
            (run / fc.FILE_NAME).write_text("repo,section,key,value\n")
            wb = openpyxl.load_workbook(fc.write_workbook(run / "summary.csv", include_for_check=False))
        self.assertEqual(wb.sheetnames, ["summary"])


class RunCsvTests(unittest.TestCase):
    def _run_dir(self, tmp: Path) -> Path:
        api = tmp / "api" / "o__r"
        git = tmp / "git" / "o__r"
        api.mkdir(parents=True)
        git.mkdir(parents=True)
        (api / "repo.json").write_text("{}")
        (git / "git_stats.json").write_text("{}")
        before = {"counts": {"prs": 5, "labels": ""}, "refs": {"refs/heads/main": "a1"}, "errors": ["x: y"]}
        after = {"counts": {"prs": 5}, "refs": {"refs/heads/main": "a1"}, "errors": []}
        fc.save_repo(api, platform="github", host="github.com", started_at="t0", before=before, after=after,
                     captured={"repo": "o/r", "merged_prs": 4})
        return tmp

    def test_rows_and_csv(self):
        with tempfile.TemporaryDirectory() as d:
            run = self._run_dir(Path(d))
            path = fc.write_run_csv(run)
            rows = list(csv.DictReader(path.open()))
        got = {(r["repo"], r["section"], r["key"]): r["value"] for r in rows}
        self.assertEqual(got[("o/r", "counts_before", "prs")], "5")
        self.assertEqual(got[("o/r", "counts_before", "labels")], "")
        self.assertEqual(got[("o/r", "ref_after", "refs/heads/main")], "a1")
        self.assertEqual(got[("o/r", "captured", "merged_prs")], "4")
        self.assertIn(("o/r", "file", "api/o__r/repo.json"), got)
        self.assertNotIn(("o/r", "file", "api/o__r/for_check.json"), got)
        self.assertEqual(got[("o/r", "error", "before")], "x: y")

    def test_workbook_has_second_sheet(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        with tempfile.TemporaryDirectory() as d:
            run = self._run_dir(Path(d))
            (run / "summary.csv").write_text("org,repo\no,r\n")
            fc.write_run_csv(run)
            xlsx = fc.write_workbook(run / "summary.csv")
            wb = openpyxl.load_workbook(xlsx)
        self.assertEqual(wb.sheetnames, ["summary", "for-check"])
        self.assertEqual([c.value for c in wb["for-check"][1]], fc.COLUMNS)

    def test_workbook_without_for_check_has_one_sheet(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl not installed")
        with tempfile.TemporaryDirectory() as d:
            (Path(d) / "summary.csv").write_text("org,repo\no,r\n")
            wb = openpyxl.load_workbook(fc.write_workbook(Path(d) / "summary.csv"))
        self.assertEqual(wb.sheetnames, ["summary"])


class LocalTargetTests(unittest.TestCase):
    def test_local_target_is_unavailable_not_an_error(self):
        import extract_org_raw_data as ex

        target = ex.RepoTarget(platform="local", org="o", full_name="o/r", meta={}, local_path=Path("."))
        snap = ex._for_check_snapshot(target, {}, "g", "l", "b", None, "github.com", "gitlab.com")
        self.assertEqual(snap["counts"], {})
        self.assertEqual(snap["errors"], [])
        self.assertIn("no platform source", snap["skipped"])


if __name__ == "__main__":
    unittest.main()

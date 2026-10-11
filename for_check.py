"""`for-check.csv`: source-side facts recorded while the platform is still reachable.

The extract alone cannot say what it *should* contain: a repo with 294 PRs and a summary row that says 12 looks
valid on its own. This module asks the platform (GitHub / GitLab API and `git ls-remote`) for counts and refs just
before and just after a repo is processed, and records what the tool itself captured, so a checker (Garmr) can
verify the extract against facts it could not recover from git history.

One long-format CSV per run, beside summary.csv:   repo,section,key,value
The same rows are the second sheet ("for-check") of summary.xlsx.

  meta          platform, host, started_at, finished_at, format_version
  counts_before entity counts before the repo was processed        (blank value = unknown, never a guess)
  counts_after  entity counts after it finished
  ref_before    branches, tags, PR/MR head refs on the source -> sha. Branch/tag names are hashed by default
                (client names / ticket ids / secrets can hide in them); see REF_MODES
  ref_after     only refs that differ from ref_before ("deleted" = gone), see meta ref_after_mode
  captured      what this tool recorded for the repo (commits, merged_prs, merged_prs_source, measured_branch,
                commits_default_branch, contributors, default_branch, size_kb)
  file          every file this run wrote under api/<repo>/ and git/<repo>/ -> size in bytes
  error         anything that could not be collected (non-fatal; collection never fails the extract)

Per-repo facts are persisted to api/<repo>/for_check.json as soon as the repo is done, so --resume runs keep the
rows of repos that are not re-processed. Metadata only: no source text is stored.
"""

from __future__ import annotations

import base64
import csv
import hashlib
import json
import logging
import os
import re
import subprocess
import urllib.parse
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from count_merged_prs import (
    github_api,
    github_headers,
    gitlab_api,
    http_get_json,
    http_post_json,
)

COLUMNS = ["repo", "section", "key", "value"]
FORMAT_VERSION = 1
FILE_NAME = "for-check.csv"
SHEET_NAME = "for-check"
REF_KEEP = re.compile(r"^refs/(heads/.+|tags/.+|pull/\d+/head|merge-requests/\d+/head)$")
# A repo with a huge ref set is skipped (and recorded as an error) rather than stalling a worker for minutes.
LS_REMOTE_TIMEOUT_SECONDS = 120
XLSX_MAX_ROWS = 1_000_000  # Excel's hard limit is 1,048,576 rows; the full detail stays in for-check.csv
# hash: branch/tag names become a stable digest (sha and ref type are kept); names: verbatim; none: no ref rows.
REF_MODES = ("hash", "names", "none")
DEFAULT_REF_MODE = "hash"
COUNT_KEYS_GITHUB = (
    "prs", "merged_prs", "issues", "labels", "milestones", "releases", "branches", "tags", "commits",
    "default_branch", "size_kb", "has_wiki", "has_issues", "issue_comments", "review_comments", "commit_comments", "collaborators",
)
COUNT_KEYS_GITLAB = (
    "prs", "merged_prs", "issues", "labels", "milestones", "releases", "branches", "tags", "commits",
    "default_branch", "collaborators", "has_wiki",
)

GITHUB_QUERY = """
query($o:String!,$r:String!){repository(owner:$o,name:$r){
  defaultBranchRef{name target{... on Commit{history{totalCount}}}}
  pullRequests{totalCount} merged:pullRequests(states:MERGED){totalCount} issues{totalCount} labels{totalCount} milestones{totalCount} releases{totalCount}
  branches:refs(refPrefix:"refs/heads/"){totalCount} tags:refs(refPrefix:"refs/tags/"){totalCount}
  diskUsage hasWikiEnabled hasIssuesEnabled}}
"""


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


_AUTH_FORMS = re.compile(r"(?i)((?:proxy-)?authorization|private-token)(['\"]?\s*[:=]\s*['\"]?)(?:(?:basic|bearer|token)\s+)?[^\s'\",;]+")
_URL_USERINFO = re.compile(r"(://)[^/\s@]+@")


def redact(text: str, *secrets: str) -> str:
    """Strip credentials from text that may land in the deliverable: the secrets themselves, their base64
    form, Authorization/PRIVATE-TOKEN header values and user:pass@ URL userinfo."""
    for secret in secrets:
        if not secret:  # "".replace("", "***") would interleave *** between every character
            continue
        text = text.replace(secret, "***").replace(base64.b64encode(secret.encode()).decode(), "***")
    text = _AUTH_FORMS.sub(lambda m: f"{m.group(1)}{m.group(2)}***", text)
    return _URL_USERINFO.sub(r"\1***@", text)


def _err(exc: BaseException) -> str:
    text = str(exc).splitlines()[0] if str(exc) else ""
    return f"{type(exc).__name__}: {text}"[:300]


def row(repo: str, section: str, key: Any, value: Any = "") -> dict[str, str]:
    return {"repo": repo, "section": section, "key": str(key), "value": "" if value is None else str(value)}


# ---------------------------------------------------------------- HTTP helpers

def _header(headers: dict[str, str], name: str) -> str:
    wanted = name.lower()
    for key, value in headers.items():
        if key.lower() == wanted:
            return value
    return ""


def _count(url: str, headers: dict[str, str]) -> int | str:
    """Item count of a list endpoint using one item per page: X-Total (GitLab) or the rel=last page (GitHub).

    "" (unknown) when neither is given and there may be more pages: GitLab omits X-Total above 10k items, and
    answering len(data)==1 then would record a 50k-item list as 1."""
    data, hdr = http_get_json(url + ("&" if "?" in url else "?") + "per_page=1", headers)
    total = _header(hdr, "X-Total")
    if total:
        return int(total)
    link = _header(hdr, "Link")
    match = re.search(r'[?&]page=(\d+)>;\s*rel="last"', link)
    if match:
        return int(match.group(1))
    if not isinstance(data, list):
        return ""
    if len(data) < 1:
        return 0
    more = 'rel="next"' in link or bool(_header(hdr, "X-Next-Page"))
    return "" if more else len(data)


def _try(errors: list[str], label: str, fn: Callable, *args: Any) -> Any:
    try:
        return fn(*args)
    except Exception as exc:  # unknown stays blank: a wrong guess would hide a real gap
        errors.append(f"{label}: {_err(exc)}")
        return ""


def _graphql_url(host: str) -> str:
    return "https://api.github.com/graphql" if host == "github.com" else f"https://{host}/api/graphql"


def _github_graphql(host: str, token: str, owner: str, repo: str) -> dict[str, Any]:
    data = http_post_json(
        _graphql_url(host),
        github_headers(token),
        {"query": GITHUB_QUERY, "variables": {"o": owner, "r": repo}},
    )
    if not isinstance(data, dict) or data.get("errors") or not data.get("data"):
        raise RuntimeError(f"graphql: {json.dumps((data or {}).get('errors', data))[:200]}")
    return data["data"]["repository"]


# ---------------------------------------------------------------- counts

def github_counts(full_name: str, *, token: str, host: str, light: bool = False) -> dict[str, Any]:
    errors: list[str] = []
    owner, _, repo = full_name.partition("/")
    counts: dict[str, Any] = {k: "" for k in COUNT_KEYS_GITHUB}
    try:
        r = _github_graphql(host, token, owner, repo)
        ref = r.get("defaultBranchRef")
        counts.update(
            prs=r["pullRequests"]["totalCount"], merged_prs=r["merged"]["totalCount"], issues=r["issues"]["totalCount"],
            labels=r["labels"]["totalCount"], milestones=r["milestones"]["totalCount"],
            releases=r["releases"]["totalCount"], branches=r["branches"]["totalCount"],
            tags=r["tags"]["totalCount"], size_kb=r.get("diskUsage", ""),
            commits=ref["target"]["history"]["totalCount"] if ref else 0,
            default_branch=ref["name"] if ref else "",
            has_wiki=str(bool(r.get("hasWikiEnabled"))).lower(),
            has_issues=str(bool(r.get("hasIssuesEnabled"))).lower(),
        )
    except Exception as exc:
        errors.append(f"graphql: {_err(exc)}")
    if light:  # the extra REST counts are only taken after the repo is processed
        counts["_errors"] = errors
        return counts
    headers = github_headers(token)
    base = f"{github_api(token, host)}/repos/{full_name}"
    counts["issue_comments"] = _try(errors, "issue_comments", _count, f"{base}/issues/comments", headers)
    counts["review_comments"] = _try(errors, "review_comments", _count, f"{base}/pulls/comments", headers)
    counts["commit_comments"] = _try(errors, "commit_comments", _count, f"{base}/comments", headers)
    counts["collaborators"] = _try(errors, "collaborators", _count, f"{base}/collaborators?affiliation=all", headers)
    counts["_errors"] = errors
    return counts


def gitlab_counts(
    project: str, *, token: str, host: str, project_id: Any = None, light: bool = False
) -> dict[str, Any]:
    errors: list[str] = []
    ident = str(project_id) if project_id else urllib.parse.quote(project, safe="")
    base = f"{gitlab_api(host)}/projects/{ident}"
    headers = {"PRIVATE-TOKEN": token, "User-Agent": "for-check"}
    counts: dict[str, Any] = {k: "" for k in COUNT_KEYS_GITLAB}
    try:
        info, _ = http_get_json(f"{base}?statistics=true", headers)
        counts["commits"] = (info.get("statistics") or {}).get("commit_count", "")
        counts["default_branch"] = info.get("default_branch") or ""
        counts["has_wiki"] = str(bool(info.get("wiki_enabled"))).lower()
    except Exception as exc:
        errors.append(f"project: {_err(exc)}")
    endpoints = (
        ("issues", "issues?scope=all"), ("prs", "merge_requests?scope=all&state=all"),
        ("merged_prs", "merge_requests?scope=all&state=merged"),
        ("branches", "repository/branches"), ("tags", "repository/tags"),
    )
    if not light:
        endpoints += (
            ("labels", "labels?include_ancestor_groups=false"), ("milestones", "milestones"),
            ("releases", "releases"), ("collaborators", "members/all"),
        )
    for key, path in endpoints:
        counts[key] = _try(errors, key, _count, f"{base}/{path}", headers)
    counts["_errors"] = errors
    return counts


# ---------------------------------------------------------------- refs

def mask_refs(refs: dict[str, str], mode: str) -> dict[str, str]:
    """Apply the ref-name privacy mode. PR/MR head refs carry only a number and are kept as they are."""
    if mode == "names":
        return dict(refs)
    if mode == "none":
        return {}
    masked: dict[str, str] = {}
    for ref, sha in refs.items():
        match = re.match(r"^(refs/(?:heads|tags)/)(.+)$", ref)
        if match:
            ref = match.group(1) + "h-" + hashlib.sha256(match.group(2).encode()).hexdigest()[:12]
        masked[ref] = sha
    return masked


def ls_remote_refs(url: str, user: str, token: str, run: Callable = subprocess.run) -> dict[str, str]:
    """Branches, tags and PR/MR head refs on the source. The token travels in the environment, never on argv."""
    auth = base64.b64encode(f"{user}:{token}".encode()).decode()
    env = {
        **os.environ, "GIT_TERMINAL_PROMPT": "0", "GIT_CONFIG_COUNT": "1",
        "GIT_CONFIG_KEY_0": "http.extraHeader", "GIT_CONFIG_VALUE_0": f"Authorization: Basic {auth}",
    }
    res = run(["git", "ls-remote", "--refs", url], capture_output=True, text=True, env=env,
              timeout=LS_REMOTE_TIMEOUT_SECONDS)
    if res.returncode != 0:
        msg = (res.stderr or "").strip().splitlines()
        raise RuntimeError(
            "git ls-remote failed: " + (redact(msg[0], token, auth) if msg else f"exit {res.returncode}")
        )
    refs: dict[str, str] = {}
    for line in res.stdout.splitlines():
        sha, _, ref = line.partition("\t")
        if ref and REF_KEEP.match(ref):
            refs[ref] = sha
    return refs


def _remote_url(platform: str, full_name: str, host: str, email: str = "") -> tuple[str, str]:
    """(git URL without credentials, basic-auth user) for `git ls-remote`."""
    if platform == "github":
        return f"https://{host}/{full_name}.git", "x-access-token"
    if platform == "gitlab":
        h = host.rstrip("/")
        scheme = "http://" if h.startswith("http://") else "https://"
        return f"{scheme}{h.replace('https://', '').replace('http://', '')}/{full_name}.git", "oauth2"
    # Repository/workspace access tokens use x-token-auth; app passwords authenticate as the account email.
    return f"https://bitbucket.org/{full_name}.git", email or "x-token-auth"


# ---------------------------------------------------------------- per-repo snapshot

def snapshot(
    platform: str, full_name: str, *, token: str, host: str, project_id: Any = None, email: str = "",
    light: bool = False, ref_mode: str = DEFAULT_REF_MODE,
) -> dict[str, Any]:
    """Counts + refs for one repo right now. Never raises: failures become `errors`.

    `light` takes the minimal call set (used for the before snapshot, to halve the API budget per repo)."""
    errors: list[str] = []
    counts: dict[str, Any] = {}
    refs: dict[str, str] = {}
    refs_ok = False
    try:
        if platform == "github":
            counts = github_counts(full_name, token=token, host=host, light=light)
        elif platform == "gitlab":
            counts = gitlab_counts(full_name, token=token, host=host, project_id=project_id, light=light)
        else:
            errors.append(f"counts: not collected for {platform}")
        errors += counts.pop("_errors", [])
    except Exception as exc:
        errors.append(f"counts: {_err(exc)}")
    if ref_mode != "none":
        try:
            url, user = _remote_url(platform, full_name, host, email)
            refs = mask_refs(ls_remote_refs(url, user, token), ref_mode)
            refs_ok = True
        except Exception as exc:
            errors.append(f"refs: {_err(exc)}")
    return {"counts": counts, "refs": refs, "refs_ok": refs_ok, "refs_mode": ref_mode, "light": light,
            "errors": [redact(e, token, email) for e in errors]}


def unavailable(reason: str) -> dict[str, Any]:
    return {"counts": {}, "refs": {}, "errors": [reason]}


def skipped(reason: str) -> dict[str, Any]:
    """Nothing to ask the platform (local/offline clone): a note, not an error."""
    return {"counts": {}, "refs": {}, "errors": [], "skipped": reason}


def _ref_diff(before: dict[str, Any], after: dict[str, Any]) -> dict[str, Any]:
    """`after` carrying only refs that changed since `before`: every PR/MR head ref used to be written twice."""
    if not (before.get("refs_ok") and after.get("refs_ok")):
        return after
    old, new = before.get("refs") or {}, after.get("refs") or {}
    diff = {k: v for k, v in new.items() if old.get(k) != v}
    diff.update({k: "deleted" for k in old if k not in new})
    return {**after, "refs": diff, "refs_mode_after": "diff_vs_before"}


def save_repo(api_dir: Path, *, platform: str, host: str, started_at: str, before: dict[str, Any],
              after: dict[str, Any], captured: dict[str, Any], log: logging.Logger | None = None) -> None:
    """Persist one repo's facts. Whatever goes wrong here must never fail the extract, but must not be silent."""
    try:
        payload = {"platform": platform, "host": host, "started_at": started_at, "finished_at": now(),
                   "before": before, "after": _ref_diff(before, after), "captured": captured}
        (api_dir / "for_check.json").write_text(json.dumps(payload, indent=2, default=str), encoding="utf-8")
    except Exception as exc:
        (log or logging.getLogger(__name__)).warning("for-check: could not save %s: %s", captured.get("repo"), _err(exc))
        try:  # keep the repo visible in for-check.csv with an error row instead of letting it vanish
            fallback = {"platform": platform, "host": host, "started_at": started_at, "finished_at": now(),
                        "captured": {k: str(v) for k, v in captured.items()}, "save_error": _err(exc)}
            (api_dir / "for_check.json").write_text(json.dumps(fallback, indent=2), encoding="utf-8")
        except Exception as exc2:
            (log or logging.getLogger(__name__)).warning("for-check: fallback save failed: %s", _err(exc2))


# ---------------------------------------------------------------- run-level CSV

def _repo_files(run_dir: Path, slug: str) -> dict[str, int]:
    files: dict[str, int] = {}
    for sub in ("api", "git"):
        base = run_dir / sub / slug
        if not base.is_dir():
            continue
        for path in sorted(base.rglob("*")):
            if path.is_file() and path.name != "for_check.json":
                files[path.relative_to(run_dir).as_posix()] = path.stat().st_size
    return files


def _current_repos(run_dir: Path) -> set[str] | None:
    """full_name of every repo in this run, so for_check.json left by a previous --resume run is not listed."""
    try:
        data = json.loads((run_dir / "repos.json").read_text(encoding="utf-8"))
        return {str(t["full_name"]) for t in data}
    except (OSError, ValueError, KeyError, TypeError):
        return None


def build_rows(run_dir: Path, repos: set[str] | None = None) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    for fc_path in sorted((run_dir / "api").glob("*/for_check.json")):
        slug = fc_path.parent.name
        try:
            info = json.loads(fc_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            rows.append(row(slug, "error", "for_check", f"unreadable {fc_path.name}: {_err(exc)}"))
            continue
        captured = info.get("captured") or {}
        label = captured.get("repo") or slug
        if repos is not None and label not in repos:
            continue
        rows.append(row(label, "meta", "format_version", FORMAT_VERSION))
        for key in ("platform", "host", "started_at", "finished_at"):
            rows.append(row(label, "meta", key, info.get(key, "")))
        if info.get("save_error"):
            rows.append(row(label, "error", "save", info["save_error"]))
        for tag in ("before", "after"):
            snap = info.get(tag) or {}
            if snap.get("skipped"):
                rows.append(row(label, "meta", f"snapshot_{tag}_skipped", snap["skipped"]))
            if snap.get("refs_mode"):
                rows.append(row(label, "meta", f"ref_{tag}_mode", snap.get("refs_mode_after") or snap["refs_mode"]))
            if snap.get("light"):
                rows.append(row(label, "meta", f"counts_{tag}_scope", "light"))
            rows += [row(label, f"counts_{tag}", k, v) for k, v in (snap.get("counts") or {}).items()]
            rows += [row(label, f"ref_{tag}", k, v) for k, v in sorted((snap.get("refs") or {}).items())]
            rows += [row(label, "error", tag, e) for e in snap.get("errors", [])]
        rows += [row(label, "captured", k, v) for k, v in captured.items() if k != "repo"]
        rows += [row(label, "file", k, v) for k, v in _repo_files(run_dir, slug).items()]
    return rows


def write_csv(path: Path, rows: list[dict[str, str]]) -> Path:
    partial = path.with_suffix(path.suffix + ".part")
    with partial.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(rows)
    partial.replace(path)
    return path


def write_run_csv(run_dir: Path) -> Path | None:
    """Rebuild <run_dir>/for-check.csv from the per-repo facts. Never raises."""
    try:
        return write_csv(run_dir / FILE_NAME, build_rows(run_dir, _current_repos(run_dir)))
    except Exception:
        return None


_NUMBER = re.compile(r"^-?(0|[1-9]\d{0,14})(\.\d{1,10})?$")
_TEXT_COLUMNS = {"org", "repo", "project_name", "default_branch", "measured_branch", "primary_language",
                 "languages_breakdown", "error", "error_class", "section", "key"}
_FORMULA_START = ("=", "+", "-", "@")


def _cell(sheet, value: str, text_only: bool):
    """Spreadsheet cell: control characters stripped (openpyxl rejects them), numbers as numbers, and text that
    a spreadsheet would execute as a formula stored as plain text."""
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if not text_only and _NUMBER.match(value):
        return float(value) if "." in value else int(value)
    cell = WriteOnlyCell(sheet, value=value)
    if value.startswith(_FORMULA_START):
        cell.data_type = "s"
    return cell


_TEXT_SECTIONS = {"meta", "error", "ref_before", "ref_after"}  # shas and names must not turn into numbers


def _append_csv(sheet, path: Path, for_check: bool, max_rows: int | None = None) -> None:
    with path.open(newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle)
        header = next(reader, [])
        sheet.append([_cell(sheet, h, True) for h in header])
        text_cols = {i for i, h in enumerate(header) if h in _TEXT_COLUMNS}
        for n, line in enumerate(reader, 1):
            if max_rows is not None and n >= max_rows:
                sheet.append([_cell(sheet, f"truncated: rows beyond {max_rows - 1} omitted, see {path.name}", True)])
                return
            text_row = for_check and len(line) > 1 and (line[1] in _TEXT_SECTIONS or line[2:3] == ["default_branch"])
            sheet.append([_cell(sheet, v, text_row or i in text_cols) for i, v in enumerate(line)])


def write_workbook(summary_csv: Path, include_for_check: bool = True) -> Path | None:
    """summary.xlsx = sheet "summary" (summary.csv) + sheet "for-check" (for-check.csv when present).

    Raises on failure: the caller logs it, so a spreadsheet problem cannot abort the run."""
    try:
        import openpyxl
    except ImportError:
        return None
    workbook = openpyxl.Workbook(write_only=True)  # streams rows, so a big for-check sheet stays out of memory
    _append_csv(workbook.create_sheet("summary"), summary_csv, for_check=False)
    fc_csv = summary_csv.with_name(FILE_NAME)
    if include_for_check and fc_csv.is_file():
        _append_csv(workbook.create_sheet(SHEET_NAME), fc_csv, for_check=True, max_rows=XLSX_MAX_ROWS)
    xlsx_path = summary_csv.with_suffix(".xlsx")
    workbook.save(xlsx_path)
    return xlsx_path
